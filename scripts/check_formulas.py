"""Recalcula no LibreOffice as fórmulas geradas pelo app e compara com calcYear(). Chamado por check_formulas.mjs."""
import json, re, subprocess, sys
from pathlib import Path
import openpyxl
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string

d = Path(sys.argv[1])
data = json.loads((d / 'in.json').read_text())

# Adaptações para o LibreOffice (só no teste): intervalo aberto Itens!E$2:E -> E$2:E3000
fix = lambda v: re.sub(r'(Itens!\$?[A-Z]+\$2:\$?[A-Z]+)(?!\$?\d)', r'\g<1>3000', v) if isinstance(v, str) else v

wb = openpyxl.Workbook(); wb.remove(wb.active)
sh = {n: wb.create_sheet(n) for n in ['Resumo', 'Gráficos', 'Itens', 'Grupos', 'Config']}
def put(rng, vals):
    s, r = rng.split('!'); col, row = coordinate_from_string(r.split(':')[0]); c0 = column_index_from_string(col)
    for i, rw in enumerate(vals):
        for j, v in enumerate(rw):
            if v != '': sh[s].cell(row + i, c0 + j, fix(v))
for blk in data['hv']: put(blk['range'], blk['values'])
put('Grupos!A2', data['rows']['grupos']); put('Itens!A2', data['rows']['itens'])
put('Config!D2', [[a] for a in data['rows']['anos']])
# FILTER não existe no LibreOffice 24.2: escreve os grupos do ano estaticamente
for k, r in enumerate([g for g in data['rows']['grupos'] if g[1] == 2026]):
    sh['Resumo'].cell(13 + k, 1, r[2]); sh['Resumo'].cell(13 + k, 2, r[3])
wb.save(d / 'in.xlsx')
subprocess.run(['soffice', '--headless', '--convert-to', 'xlsx:Calc MS Excel 2007 XML', '--outdir', str(d / 'out'), str(d / 'in.xlsx')], check=True, capture_output=True)

R = openpyxl.load_workbook(d / 'out' / 'in.xlsx', data_only=True)['Resumo']
c, ok = data['c'], True
for r, k in {4: 'receitas', 5: 'meta', 6: 'recMenosInv', 7: 'despesas', 8: 'invest', 9: 'gastoReal', 10: 'saldo'}.items():
    got = [round(R.cell(r, col).value or 0, 2) for col in range(3, 15)]
    if got != c[k]: ok = False; print(f'ERRO linha {r} ({k}): planilha={got} app={c[k]}')
for gi, g in enumerate(c['grupos']):
    got = [round(R.cell(14 + gi, col).value or 0, 2) for col in range(3, 15)]
    if got != g['tot']: ok = False; print(f"ERRO grupo {g['g']['nome']}")
if R['B2'].value != data['n']: ok = False; print('ERRO meses com lançamento', R['B2'].value, data['n'])
print('Fórmulas OK: Resumo bate com o app' if ok else 'Fórmulas com divergência')
sys.exit(0 if ok else 1)
